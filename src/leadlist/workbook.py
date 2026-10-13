"""Writing and reading Excel workbooks."""
from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E6E6E6")
_SORT_ARROWS = ("↑", "↓")

# A cell that opens with one of these is run as a formula by Excel, LibreOffice or Numbers.
# The line feed is here because a carriage return comes back as one when the file is read: XML has only the one.
FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r", "\n")
_QUOTE = "'"
# C0 and C1 control characters, except tab, line feed and carriage return. openpyxl refuses some of them.
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")
# Half of a pair on its own. It cannot be written as UTF-8, so no file could hold it.
_LONE_SURROGATE = re.compile(r"[\ud800-\udfff]")


def clean_text(text: str) -> str:
    """Drop control characters and replace lone surrogates: what a hostile page can send and no file can hold."""
    return _CONTROL.sub("", _LONE_SURROGATE.sub("�", text))


def neutralise(text: str) -> str:
    """Put a quote in front of text a spreadsheet would run as a formula, so that it stays text."""
    # A value that already opens with quotes and then a trigger gets one more,
    # so that read_sheet can take exactly one off again and nothing is changed by the round trip.
    return _QUOTE + text if text.lstrip(_QUOTE).startswith(FORMULA_TRIGGERS) else text


def restore(text: str) -> str:
    """Undo neutralise."""
    return text[1:] if text.startswith(_QUOTE) and text.lstrip(_QUOTE).startswith(FORMULA_TRIGGERS) else text


def _cell_text(value: object) -> object:
    return neutralise(clean_text(value)) if isinstance(value, str) else value


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> Path:
    """One sheet per frame, with a frozen bold header, a filter and sized columns.

    Names in these sheets come from scraped pages and CRM exports. Every text
    cell, headers included, is cleaned and kept from being read as a formula.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            safe = frame.apply(lambda column: column.map(_cell_text))
            safe.columns = [_cell_text(column) for column in frame.columns]
            safe.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            for row in sheet.iter_rows():
                for cell in row:
                    # Should the cleaning above ever miss one, it is still stored as text here.
                    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith(FORMULA_TRIGGERS)):
                        cell.value = _QUOTE + str(cell.value)
                        cell.data_type = "s"
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = _HEADER_FONT
                cell.fill = _HEADER_FILL
            if sheet.max_column:
                sheet.auto_filter.ref = sheet.dimensions
            for index, column in enumerate(safe.columns, start=1):
                longest = max([len(str(column))] + [len(str(v)) for v in safe[column].head(500)])
                sheet.column_dimensions[get_column_letter(index)].width = min(60, longest + 2)
    return path


def read_sheet(path: Path, sheet: str) -> pd.DataFrame:
    """A sheet this program wrote: header in row one, every value as text."""
    # na_filter off: a person or organisation called "NA" or "None" is a name, not a missing value.
    frame = pd.read_excel(path, sheet_name=sheet, dtype=str, na_filter=False).fillna("")
    # Take off the quote write_workbook put on, so that a name is the same going out and coming back.
    # A sheet edited by hand to open a cell with a quote and then a formula character loses that quote here.
    frame = frame.apply(lambda column: column.map(lambda value: restore(value) if isinstance(value, str) else value))
    frame.columns = [restore(column) if isinstance(column, str) else column for column in frame.columns]
    return frame


def _label(value: object) -> str:
    """A header cell of a report export as text, without the arrow that marks the sorted column."""
    # Plain string work. A pattern for "spaces, arrow, spaces, end" takes quadratic time on a long run of spaces.
    text = str(value).strip()
    return text[:-1].rstrip() if text.endswith(_SORT_ARROWS) else text


def read_report(path: Path, must_have: list[str]) -> pd.DataFrame:
    """Read a CRM report export, wherever its header row happens to sit.

    Report exports open with a title, filter notes and blank rows, may leave
    the first column empty, and mark the sorted column with an arrow. The
    header is the first row that holds every column in `must_have`.

    The export is someone else's file: nothing is stripped from its values.
    """
    raw = pd.read_excel(path, header=None, dtype=object)
    for index, row in raw.iterrows():
        labels = ["" if pd.isna(v) else _label(v) for v in row]
        if all(name in labels for name in must_have):
            body = raw.iloc[index + 1:].copy()
            body.columns = labels
            body = body.loc[:, [label != "" for label in labels]]
            return body.dropna(how="all").reset_index(drop=True)
    raise SystemExit(f"{path.name}: found no header row with the columns {must_have}")
